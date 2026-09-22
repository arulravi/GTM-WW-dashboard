#!/usr/bin/env python3
"""
Opex & HC Outlook dashboard -- data refresh (ETL)
=================================================
Pulls straight from the finance SQL views and rebuilds `data.js`, which powers
`dashboard.html` (a self-contained, offline executive dashboard).

    python refresh.py

What it does
------------
* Queries vw_Expense_Final_LDAP_Filtered and vw_Headcount_Final_LDAP_Filtered
  (Windows auth) at CCH L2/L3/L4 granularity.
* Normalises versions to friendly keys (Actuals / Plan / QRF / Forecast) and
  excludes sales Commissions -- reconciled to the WK-OL workbook.
* Collapses each quarter's monthly headcount using the "current actual month"
  convention (matches the workbook).
* Preserves commentary typed in the dashboard (merges commentary.json).

Output: data.js  (assigns window.OPEX_DATA)
"""
from __future__ import annotations

import json
import os
import re
import sys
import datetime

import openpyxl

import db

# When packaged as an .exe (PyInstaller), read/write next to the exe so the
# refreshed data.js and commentary.json live beside the app, not in a temp dir.
HERE = (os.path.dirname(sys.executable) if getattr(sys, "frozen", False)
        else os.path.dirname(os.path.abspath(__file__)))
FISCAL_YEARS = ["2024", "2025", "2026"]


def log(*a):
    print("[refresh]", *a, flush=True)


def _i(v) -> int:
    try:
        if v is None or (isinstance(v, float) and v != v):  # NaN
            return 0
        return int(round(float(v)))
    except (TypeError, ValueError):
        return 0


def _f(v, nd=1) -> float:
    try:
        if v is None or (isinstance(v, float) and v != v):
            return 0.0
        return round(float(v), nd)
    except (TypeError, ValueError):
        return 0.0


def q_key(qtr: str) -> tuple[int, int]:
    m = re.match(r"(\d{4})-Q(\d)", qtr or "")
    return (int(m.group(1)), int(m.group(2))) if m else (0, 0)


def norm_version(ver: str, current_q_num: str) -> str | None:
    """Map a raw Version string to a stable dashboard key."""
    if ver == "Actuals":
        return "act"
    if ver == "Forecast":
        return "fcst"
    m = re.search(r"FY\d\d\s+Plan", ver) or re.match(r"FY\d\d Plan", ver)
    if "Plan" in ver:
        return "plan"
    mq = re.search(r"Q(\d)\s*QRF", ver)
    if mq:
        return f"qrf{mq.group(1)}"
    mw = re.search(r"Outlook\s+WK\s*(\d+)", ver, re.IGNORECASE)
    if mw:
        return f"wk{mw.group(1)}"
    return None


def read_manual_inputs() -> dict:
    """Read the analyst-maintained 'Manual Inputs.xlsx' (Open Reqs, Severance,
    Salary) if present. Returns {'open_reqs': {'l3|l4|qtr': n},
    'severance': {'l3|qtr': m}, 'salary': {'position_id': annual_$}}."""
    out = {"open_reqs": {}, "severance": {}, "salary": {}}
    path = os.path.join(HERE, "Manual Inputs.xlsx")
    if not os.path.isfile(path):
        return out
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        if "Open Reqs" in wb.sheetnames:
            for row in wb["Open Reqs"].iter_rows(min_row=6, values_only=True):
                l3, l4, q, v = (row + (None,) * 5)[1:5]
                if l3 and l4 and q and v not in (None, ""):
                    out["open_reqs"][f"{l3}|{l4}|{q}"] = _f(v)
        if "Severance" in wb.sheetnames:
            for row in wb["Severance"].iter_rows(min_row=6, values_only=True):
                l3, q, v = (row + (None,) * 4)[1:4]
                if l3 and q and v not in (None, ""):
                    out["severance"][f"{l3}|{q}"] = _f(v, 3)
        if "Salary" in wb.sheetnames:
            for row in wb["Salary"].iter_rows(min_row=6, values_only=True):
                pid, _title, v = (row + (None,) * 4)[1:4]
                if pid and v not in (None, ""):
                    out["salary"][str(pid)] = _f(v, 2)
        wb.close()
        log(f"  merged Manual Inputs.xlsx (open_reqs {len(out['open_reqs'])}, "
            f"severance {len(out['severance'])}, salary {len(out['salary'])})")
    except Exception as e:
        log("  WARN Manual Inputs.xlsx:", e)
    return out


SENTINEL_DATE = "1900-01-01"  # the view's null-date placeholder, not a real date


def _date_str(v) -> str | None:
    if v is None:
        return None
    s = str(v)
    if s in ("None", "NaT", "nan", ""):
        return None
    s = s[:10]
    return None if s == SENTINEL_DATE else s


def build_position_analysis(current_q: str | None, cq_num: str) -> list:
    """
    Position/req-level detail for the Analysis tab: for every position, its
    Start Date and Planned Term Date under each version (Forecast, QRF, any
    WK-Outlook snapshot). The dashboard diffs the same Position ID across two
    versions (e.g. WK7 OL vs Forecast) to flag hires pushed out/pulled in and
    terminations whose timing moved -- the SQL view carries several rows per
    (Position ID, Version) for different pipeline stages, so those are
    collapsed here into one row per (Position ID, Version).
    """
    if not current_q:
        return []
    raw = db.position_snapshot(current_q)
    if raw.empty:
        return []
    # A Position ID can carry multiple distinct rows within the same version
    # (e.g. the past incumbent's actual historical hire date AND a separate
    # open-backfill req's projected date). Sort pipeline-flagged rows first so
    # the "first" aggregation below prefers the pipeline row's Start Date
    # when one exists, instead of an unrelated old actual-hire date.
    is_pipe = (raw["committed"].fillna(0) != 0) | (raw["open_reqs"].fillna(0) != 0) \
        | (raw["pub"].fillna(0) != 0) | (raw["unpub"].fillna(0) != 0)
    raw = raw.assign(_pipe=is_pipe.astype(int)).sort_values("_pipe", ascending=False)
    g = raw.groupby(["pid", "ver"], as_index=False).agg(
        req=("req", "first"), title=("title", "first"), pdesc=("pdesc", "first"),
        l3=("l3", "first"), l4=("l4", "first"), tbh=("tbh", "first"),
        start_dt=("start_dt", "first"), term_dt=("term_dt", "first"),
        committed=("committed", "sum"), open_reqs=("open_reqs", "sum"),
        pub=("pub", "sum"), unpub=("unpub", "sum"),
    )
    out = []
    for r in g.itertuples(index=False):
        vk = norm_version(r.ver, cq_num)
        if vk is None:
            continue
        pipeline = 1 if (_f(r.committed) + _f(r.open_reqs) + _f(r.pub) + _f(r.unpub)) != 0 else 0
        out.append([
            r.pid, r.req or "", r.title or "", r.pdesc or "",
            r.l3 or "(none)", r.l4 or "(none)", vk,
            _date_str(r.start_dt), _date_str(r.term_dt), pipeline, r.tbh or "",
        ])
    return out


def build_pipeline_positions(current_q: str | None) -> list:
    """
    Row-level detail behind the Monthly Hiring pipeline (Committed/Open Reqs/
    Published/Unpublished) -- powers the drill-through from a forecast cell to
    the actual positions/reqs summing to it. Only current + future quarters
    (completed quarters use actual hires, which have no pipeline detail).
    Bucketed by the position's own Start Date month, same as
    db.hc_monthly_pipeline's aggregate (so a drill-through's rows always sum
    to the cell that was clicked).
    """
    if not current_q:
        return []
    raw = db.position_pipeline_rows(FISCAL_YEARS)
    if raw.empty:
        return []
    PIPE_KINDS = [("committed", "committed"), ("open_reqs", "open_reqs"),
                  ("published", "published"), ("unpublished", "unpublished")]
    out = []
    for r in raw.itertuples(index=False):
        if q_key(r.qtr) < q_key(current_q):
            continue
        start_month = _date_str(r.planned_start)
        if not start_month:
            continue
        start_month = start_month[:7]
        for attr, kind in PIPE_KINDS:
            if _i(getattr(r, attr)) == 0:
                continue
            out.append({
                "l3": r.l3 or "(none)", "l4": r.l4 or "(none)", "qtr": r.qtr,
                "month": start_month, "kind": kind,
                "pid": r.pid, "req": r.req or "",
                "job_family": r.job_family or "", "job_level": r.job_level or "",
                "created": _date_str(r.created_dt),
                "status_stage": r.status_stage or "", "status_master": r.status_master or "",
                "status_wd": r.status_wd or "",
                "work_city": r.work_city or "", "pdesc": r.pdesc or "",
                "manager": r.manager or "", "candidate": r.candidate or "",
                "univ_cohort": r.univ_cohort or "",
                "planned_start": _date_str(r.planned_start),
                "req_hire_date": _date_str(r.req_hire_date),
                "planned_term": _date_str(r.planned_term),
                "func_controller": r.func_controller or "",
            })
    return out


SNAPSHOTS_SEEN_PATH = os.path.join(HERE, "snapshots_seen.json")


def rollover_snapshot_commentary(commentary: dict, wk_keys: list[str],
                                 current_qrf: str) -> bool:
    """Carry manual commentary forward automatically when a new Outlook (WK)
    snapshot is taken.

    Commentary is keyed by the (Current, vs-Target) pair -- e.g. the notes the
    team writes for "Forecast vs WK7 OL" are stored under `<field>__fcst_vs_wk7`
    (this covers the per-cost-element Commentary, Additional Notes, and Risks &
    Opportunities alike -- they're all just fields in the same store).

    When Finance publishes a new snapshot (say WK10 OL), the Forecast that was
    live at that moment IS what got frozen into WK10. So every "Forecast vs X"
    note should now also read as "WK10 vs X". This function detects newly
    appeared WK snapshots and copies each `<field>__fcst_vs_<X>` note (and the
    plain default "Forecast vs current-QRF" note) to `<field>__<newWk>_vs_<X>`,
    so the exact same commentary follows the same comparison into the new cycle
    with no re-typing. It never overwrites a note that already exists for the
    new pairing, so anything a user has already edited there is left untouched.

    A `snapshots_seen.json` marker records which snapshots have been processed,
    so the rollover fires exactly once per new snapshot. On the very first run
    (no marker yet) it only records a baseline and rolls nothing retroactively.

    Returns True if any commentary was added (so the caller re-persists it).
    """
    first_run = not os.path.isfile(SNAPSHOTS_SEEN_PATH)
    seen: list[str] = []
    if not first_run:
        try:
            with open(SNAPSHOTS_SEEN_PATH, "r", encoding="utf-8") as f:
                seen = json.load(f)
        except Exception:
            seen = []

    current = sorted({k for k in wk_keys if re.fullmatch(r"wk\d+", k)},
                     key=lambda k: int(k[2:]))

    def _save_seen():
        try:
            with open(SNAPSHOTS_SEEN_PATH, "w", encoding="utf-8") as f:
                json.dump(current, f)
        except Exception as e:
            log("  WARN could not write snapshots_seen.json:", e)

    if first_run:
        _save_seen()
        log(f"  snapshot rollover: baseline recorded ({', '.join(current) or 'none'})")
        return False

    new_snaps = [k for k in current if k not in seen]
    if not new_snaps:
        return False

    added = 0
    for scope, notes in commentary.items():
        if not isinstance(notes, dict):
            continue
        extra: dict[str, str] = {}
        for key, val in notes.items():
            if not isinstance(val, str) or not val.strip():
                continue
            m = re.match(r"^(.*)__fcst_vs_(.+)$", key)
            if m:
                field, tgt = m.group(1), m.group(2)
            elif "__" not in key:            # plain default pairing = Forecast vs current QRF
                field, tgt = key, current_qrf
            else:
                continue                     # already a snapshot-vs-something pairing; leave as-is
            is_wk = lambda s: re.fullmatch(r"wk\d+", s) is not None
            for wk in new_snaps:
                # forward: the new snapshot vs the prior target (e.g. WK12 vs WK10)
                pairs = [f"{field}__{wk}_vs_{tgt}"]
                # reverse: prior snapshot vs the new one (WK10 vs WK12) -- only
                # when the target is itself a WK snapshot, since both are then
                # selectable as Current/vs-Target, so the note shows whichever
                # way the user picks. (No reverse for Plan/QRF targets: those
                # can't be chosen as "Current".)
                if is_wk(tgt):
                    pairs.append(f"{field}__{tgt}_vs_{wk}")
                for nk in pairs:
                    if nk not in notes and nk not in extra:
                        extra[nk] = val
                        added += 1
        notes.update(extra)

    _save_seen()
    log(f"  snapshot rollover: new {', '.join(new_snaps)} -> carried {added} commentary entries forward")
    return added > 0


def build_data() -> dict:
    log("Connecting and pulling opex facts…")
    opex = db.opex_facts(FISCAL_YEARS)
    log(f"  opex rows: {len(opex):,}")
    hc = db.hc_facts(FISCAL_YEARS)
    log(f"  hc rows: {len(hc):,}")

    current_q = db.current_quarter()
    cq_num = re.search(r"Q(\d)", current_q or "").group(1) if current_q else "3"
    log(f"  current quarter: {current_q}")

    plan_fy = current_q.split("-")[0] if current_q else "2026"
    planning = db.hc_planning(FISCAL_YEARS)
    log(f"  hc planning rows: {len(planning):,}")
    position_out = build_position_analysis(current_q, cq_num)
    log(f"  position analysis rows: {len(position_out):,}")
    pipeline_positions_out = build_pipeline_positions(current_q)
    log(f"  pipeline position rows: {len(pipeline_positions_out):,}")

    # ---- opex facts -> compact array [l2,l3,l4,ce,qtr,verKey,val] ----------
    opex_out = []
    hierarchy: dict = {}
    for r in opex.itertuples(index=False):
        vk = norm_version(r.ver, cq_num)
        if vk is None:
            continue
        l2, l3, l4 = r.l2 or "(none)", r.l3 or "(none)", r.l4 or "(none)"
        opex_out.append([l2, l3, l4, r.ce, r.qtr, vk, round(float(r.v or 0), 3)])
        hierarchy.setdefault(l2, {}).setdefault(l3, set()).add(l4)

    hier_json = {l2: {l3: sorted(l4s) for l3, l4s in l3s.items()}
                 for l2, l3s in hierarchy.items()}

    # SQL stops publishing a 'Forecast' version once a quarter closes --
    # confirmed directly: 2026-Q3 has 0 Forecast rows (only Actuals), while
    # 2026-Q4 (the live quarter) has ~22K. Standard FP&A convention once a
    # quarter is closed is that Forecast = Actuals (the books are closed,
    # there's nothing left to forecast) -- without this, "Current: Forecast"
    # silently shows $0 for every closed quarter instead of the real number.
    # Backfill a synthetic "fcst" row from each "act" row for any quarter
    # that has Actuals but no Forecast at all.
    _fcst_qtrs = {row[4] for row in opex_out if row[5] == "fcst"}
    opex_out += [[r[0], r[1], r[2], r[3], r[4], "fcst", r[6]]
                 for r in opex_out if r[5] == "act" and r[4] not in _fcst_qtrs]

    # ---- GL-level detail for the drill-down: [l3,l4,ce,gl,qtr,verKey,val] ---
    # only years the walk shows (FY25+), to keep the file lean
    gl = db.opex_gl([y for y in FISCAL_YEARS if y >= "2025"])
    log(f"  gl rows: {len(gl):,}")
    gl_out = []
    for r in gl.itertuples(index=False):
        vk = norm_version(r.ver, cq_num)
        if vk is None:
            continue
        val = round(float(r.v or 0), 3)
        if abs(val) < 0.005:            # drop noise to keep the file lean
            continue
        gl_out.append([r.l3 or "(none)", r.l4 or "(none)", r.ce,
                       (r.gl or "(none)"), r.qtr, vk, val])

    # Same Forecast-for-closed-quarters backfill as opex_out above -- keeps
    # the GL drill-down tying to the walk table instead of showing $0 rows
    # for any cost element viewed under "Current: Forecast" on a closed qtr.
    _gl_fcst_qtrs = {row[4] for row in gl_out if row[5] == "fcst"}
    gl_out += [[r[0], r[1], r[2], r[3], r[4], "fcst", r[6]]
               for r in gl_out if r[5] == "act" and r[4] not in _gl_fcst_qtrs]

    # ---- headcount: collapse each (scope, qtr, ver) to one value -----------
    # "current actual month" = 2nd-to-last period when a forecast tail exists.
    hc_group: dict = {}
    for r in hc.itertuples(index=False):
        vk = norm_version(r.ver, cq_num)
        if vk is None:
            continue
        key = (r.l2 or "(none)", r.l3 or "(none)", r.l4 or "(none)", r.qtr, vk)
        hc_group.setdefault(key, {})[r.period] = float(r.hc or 0)

    hc_out = []
    for (l2, l3, l4, qtr, vk), periods in hc_group.items():
        if vk == "act":
            continue  # Actuals HC double-counts; excluded (see README)
        ordered = [periods[p] for p in sorted(periods)]
        if not ordered:
            continue
        val = ordered[-2] if len(ordered) >= 3 else ordered[-1]
        hc_out.append([l2, l3, l4, qtr, vk, round(val, 1)])

    # ---- actual (Active) headcount -> quarter-ending value, version 'act' ---
    hc_act = db.hc_actuals_ending(FISCAL_YEARS)
    act_group: dict = {}
    for r in hc_act.itertuples(index=False):
        key = (r.l2 or "(none)", r.l3 or "(none)", r.l4 or "(none)", r.qtr)
        act_group.setdefault(key, {})[r.period] = float(r.hc or 0)
    for (l2, l3, l4, qtr), periods in act_group.items():
        if not periods:
            continue
        ending = periods[sorted(periods)[-1]]   # last (most recent) actual period
        hc_out.append([l2, l3, l4, qtr, "act", round(ending, 1)])

    # Same Forecast-for-closed-quarters backfill as opex_out above (row shape
    # here is [l2,l3,l4,qtr,vk,val] -- qtr/vk/val sit at different indices
    # than opex_out's 7-column rows).
    _hc_fcst_qtrs = {row[3] for row in hc_out if row[4] == "fcst"}
    hc_out += [[r[0], r[1], r[2], r[3], "fcst", r[5]]
               for r in hc_out if r[4] == "act" and r[3] not in _hc_fcst_qtrs]

    # ---- HC by role family (current quarter) -> array ---------------------
    def _fam(desc):
        if not desc or not str(desc).strip():
            return "OTHER"
        f = re.split(r"\s[-–]\s", str(desc).strip())[0].strip().upper()
        return f or "OTHER"
    def _flabel(f):
        return f if len(f) <= 4 else f.title()   # keep acronyms (SC, QBSR) upper
    role_raw = db.hc_by_role(current_q)
    cellr = {}       # (l3,l4,vk,fam,period) -> hc
    periodsr = {}    # (l3,l4,vk) -> set(periods)
    for r in role_raw.itertuples(index=False):
        vk = norm_version(r.ver, cq_num)
        if vk is None:
            continue
        l3 = r.l3 or "(none)"; l4 = r.l4 or "(none)"; fam = _fam(r.role)
        cellr[(l3, l4, vk, fam, r.period)] = cellr.get((l3, l4, vk, fam, r.period), 0.0) + float(r.hc or 0)
        periodsr.setdefault((l3, l4, vk), set()).add(r.period)
    agg = {}
    for (l3, l4, vk, fam, period), h in cellr.items():
        ps = sorted(periodsr[(l3, l4, vk)])
        rep = ps[-2] if len(ps) >= 3 else ps[-1]   # same "current actual month" rule
        if period == rep:
            agg[(l3, l4, vk, fam)] = agg.get((l3, l4, vk, fam), 0.0) + h
    role_out = [[l3, l4, _flabel(fam), vk, round(h, 1)]
                for (l3, l4, vk, fam), h in agg.items() if round(h) != 0]

    # ---- HC by role family, year-end: FY25 (2025-Q4 Actuals) vs FY26
    #      (2026-Q4 Forecast) -> [l3, l4, famLabel, 'fy25'|'fy26', hc]. Takes
    #      each (l3,l4,role)'s LAST fiscal period as the year-end point-in-time
    #      headcount. Drives the "FY25 vs FY26 HC Role Walk" chart. ------------
    def _role_yearend(df):
        cell, periods = {}, {}
        for r in df.itertuples(index=False):
            key = (r.l3 or "(none)", r.l4 or "(none)", _fam(r.role))
            cell[(key, r.period)] = cell.get((key, r.period), 0.0) + float(r.hc or 0)
            periods.setdefault(key, set()).add(r.period)
        out = {}
        for (key, period), h in cell.items():
            if period == max(periods[key]):        # ending period = year-end
                out[key] = out.get(key, 0.0) + h
        return out
    role_year_out = []
    try:
        fy25 = _role_yearend(db.hc_role_ending("2025-Q4", "Actuals", active_only=True))
        fy26 = _role_yearend(db.hc_role_ending("2026-Q4", "Forecast"))
        for (l3, l4, fam), h in fy25.items():
            if round(h): role_year_out.append([l3, l4, _flabel(fam), "fy25", round(h, 1)])
        for (l3, l4, fam), h in fy26.items():
            if round(h): role_year_out.append([l3, l4, _flabel(fam), "fy26", round(h, 1)])
        log(f"  role-year (FY25 vs FY26) rows: {len(role_year_out):,}")
    except Exception as e:
        log("  WARN role-year build:", e)

    # ---- Monthly hiring: actual (completed quarters) + pipeline (current/future) ---
    # Actual: `Actual Hires` posts once per quarter on a single fiscal period; find
    # that "snapshot period" per quarter (the one with the largest total across
    # every org, confirmed uniform in validation), then keep only its rows so a
    # quarter's hires aren't counted twice across periods. Bucketed by each
    # hire's own Req Hire Date -> true calendar month.
    # Prior + current fiscal year, so Monthly Hiring can show a full
    # year-over-year trend instead of just the current fiscal year.
    mh_raw = db.hc_monthly_hires([str(int(plan_fy) - 1), plan_fy])
    qtr_period_totals: dict = {}
    for r in mh_raw.itertuples(index=False):
        qtr_period_totals[(r.qtr, r.period)] = qtr_period_totals.get((r.qtr, r.period), 0.0) + float(r.hires or 0)
    snapshot_period = {}
    for (qtr, period), tot in qtr_period_totals.items():
        if qtr not in snapshot_period or tot > snapshot_period[qtr][1]:
            snapshot_period[qtr] = (period, tot)
    monthly_hires_out = []
    for r in mh_raw.itertuples(index=False):
        if not r.hire_month:
            continue
        sp = snapshot_period.get(r.qtr)
        if not sp or r.period != sp[0]:
            continue
        val = _i(r.hires)
        if val == 0:
            continue
        monthly_hires_out.append([r.l3 or "(none)", r.l4 or "(none)", r.hire_month, r.qtr, "act", val])

    # Pipeline: for the current + future quarters, incremental positions (not yet
    # hired) bucketed by each position's own expected Start Date -> calendar month.
    pipe_raw = db.hc_monthly_pipeline(FISCAL_YEARS)
    PIPE_KINDS = [("committed", "committed"), ("open_reqs", "open_reqs"),
                  ("published", "published"), ("unpublished", "unpublished")]
    for r in pipe_raw.itertuples(index=False):
        if not r.start_month or q_key(r.qtr) < q_key(current_q or ""):
            continue   # only current + future quarters -- completed quarters use actuals above
        for attr, kind in PIPE_KINDS:
            val = _i(getattr(r, attr))
            if val == 0:
                continue
            monthly_hires_out.append([r.l3 or "(none)", r.l4 or "(none)", r.start_month, r.qtr, kind, val])
    log(f"  monthly hires: {len(monthly_hires_out):,} rows, snapshot periods {snapshot_period}")

    # ---- Trended HC walk: movement drivers per L3 x L4 x quarter ----------
    mov = db.hc_movement(FISCAL_YEARS)
    FKEYS = ["beg", "hires", "committed", "open_reqs", "pub", "unpub", "abf",
             "iadj", "imo", "tin", "tout", "terms", "known", "attr",
             "mgmt_hedge", "mgmt_tin", "mgmt_tout", "fcst_adj", "endhc"]
    mov_idx = {}
    l4s_by_l3 = {}
    for r in mov.itertuples(index=False):
        l4 = r.l4 or "(none)"
        mov_idx[(r.l3, l4, r.qtr, r.ver)] = r
        l4s_by_l3.setdefault(r.l3, set()).add(l4)
    walk_qs = sorted({r.qtr for r in mov.itertuples(index=False)}, key=q_key)
    walk_cells, walk_vertype = [], {}
    for qtr in walk_qs:
        vt = "Actuals" if q_key(qtr) < q_key(current_q or "") else "Forecast"
        walk_vertype[qtr] = vt
        for l3 in db.config.SCOPE_L3:
            # "__all" is the L3-wide total, built by summing every L4 under it,
            # so the walk with no L4 filter matches the old L3-only behavior.
            totals = {fk: 0.0 for fk in FKEYS}
            for l4 in l4s_by_l3.get(l3, ()):
                r = mov_idx.get((l3, l4, qtr, vt)) or mov_idx.get((l3, l4, qtr, "Actuals")) \
                    or mov_idx.get((l3, l4, qtr, "Forecast"))
                if r is None:
                    continue
                for fk in FKEYS:
                    v = _f(getattr(r, fk), 0)
                    walk_cells.append([l3, l4, qtr, fk, v])
                    totals[fk] += v
            for fk in FKEYS:
                walk_cells.append([l3, "__all", qtr, fk, totals[fk]])
    hcwalk = {"quarters": walk_qs, "vertype": walk_vertype, "cells": walk_cells}

    # ---- HC planning -> array ---------------------------------------------
    plan_out = []
    for r in planning.itertuples(index=False):
        vk = norm_version(r.ver, cq_num)
        if vk is None:
            continue
        plan_out.append({
            "l2": r.l2 or "(none)", "l3": r.l3 or "(none)", "l4": r.l4 or "(none)",
            "role_l2": getattr(r, "role_l2", None) or "", "role_l3": getattr(r, "role_l3", None) or "",
            "qtr": r.qtr, "ver": vk,
            "hires": _i(r.hires), "terms": _i(r.terms),
            "known_terms": _i(r.known_terms), "open_reqs": _i(r.open_reqs),
            "attrition": _f(r.attrition), "committed": _i(r.committed),
            "tin": _i(r.tin), "tout": _i(r.tout),
            "internal_adj": _f(r.internal_adj), "mgmt_hedge": _f(r.mgmt_hedge),
            "forecast_adj": _f(r.forecast_adj),
            "beg_hc": _f(r.beg_hc), "end_hc": _f(r.end_hc),
        })

    # Same Forecast-for-closed-quarters backfill as opex_out above --
    # confirmed the same gap here too (2026-Q3 has Actuals but 0 Forecast
    # rows in this planning view). Keeps the Headcount Walk/HC KPIs from
    # showing 0 for any closed quarter under "Current: Forecast".
    _plan_fcst_keys = {(p["l2"], p["l3"], p["l4"], p["qtr"]) for p in plan_out if p["ver"] == "fcst"}
    plan_out += [dict(p, ver="fcst") for p in plan_out
                 if p["ver"] == "act" and (p["l2"], p["l3"], p["l4"], p["qtr"]) not in _plan_fcst_keys]

    # ---- available versions (for the dropdowns), in a sensible order -------
    present = {row[5] for row in opex_out}
    wk_keys = sorted((k for k in present if k.startswith("wk")), key=lambda k: int(k[2:]))
    order = ["act", "plan", "qrf1", "qrf2", "qrf3", "qrf4"] + wk_keys + ["fcst"]
    labels = {"act": "Actuals", "plan": f"FY{plan_fy[-2:]} Plan", "fcst": "Forecast",
              "qrf1": "Q1 QRF", "qrf2": "Q2 QRF", "qrf3": "Q3 QRF", "qrf4": "Q4 QRF"}
    labels.update({k: f"WK {k[2:]} OL" for k in wk_keys})
    versions = [{"key": k, "label": labels[k]} for k in order if k in present]

    quarters = sorted({row[4] for row in opex_out}, key=q_key)

    # ---- merge saved commentary -------------------------------------------
    commentary = {}
    cpath = os.path.join(HERE, "commentary.json")
    if os.path.isfile(cpath):
        try:
            with open(cpath, "r", encoding="utf-8") as f:
                commentary = json.load(f)
            log("  merged commentary.json")
        except Exception as e:
            log("  WARN commentary.json:", e)

    # Deletion-approval metadata (see server.py's Add Commentary Deletion
    # Approval Control) -- baked in too so a fresh page load (before the
    # client's own live poll of /api/commentary lands) already knows which
    # fields are protected, and so a Save Final Snapshot export carries the
    # same protection state as the live app.
    commentary_meta = {}
    mpath = os.path.join(HERE, "commentary_meta.json")
    if os.path.isfile(mpath):
        try:
            with open(mpath, "r", encoding="utf-8") as f:
                commentary_meta = json.load(f)
        except Exception as e:
            log("  WARN commentary_meta.json:", e)

    deletion_requests = []
    dpath = os.path.join(HERE, "deletion_requests.json")
    if os.path.isfile(dpath):
        try:
            with open(dpath, "r", encoding="utf-8") as f:
                deletion_requests = json.load(f)
        except Exception as e:
            log("  WARN deletion_requests.json:", e)

    # Auto-carry manual commentary forward when a new Outlook snapshot appears
    # (Forecast vs WKx notes become WKnew vs WKx, etc.). Non-destructive; when
    # it adds anything, persist commentary.json so the rollover survives.
    try:
        if rollover_snapshot_commentary(commentary, wk_keys, f"qrf{cq_num}"):
            with open(cpath, "w", encoding="utf-8") as f:
                json.dump(commentary, f, ensure_ascii=False, indent=2)
            log("  wrote rolled-over commentary.json")
    except Exception as e:
        log("  WARN snapshot commentary rollover:", e)

    data = {
        "meta": {
            "generated_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
            "current_quarter": current_q,
            "current_qrf": f"qrf{cq_num}",
            "plan_fy": plan_fy,
            "source": f"{db.config.SQL_SERVER} · {db.config.SQL_DATABASE}",
        },
        "ce_order": db.config.COST_ELEMENT_ORDER,
        "hierarchy": hier_json,
        "versions": versions,
        "quarters": quarters,
        "opex": opex_out,
        "gl": gl_out,
        "hc": hc_out,
        "role": role_out,
        "role_year": role_year_out,
        "monthly_hires": monthly_hires_out,
        "hcwalk": hcwalk,
        "planning": plan_out,
        "positions": position_out,
        "pipeline_positions": pipeline_positions_out,
        "manual": read_manual_inputs(),
        "commentary": commentary,
        "commentary_meta": commentary_meta,
        "deletion_requests": deletion_requests,
    }
    return data


def write_data_js(data: dict) -> str:
    out = os.path.join(HERE, "data.js")
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    with open(out, "w", encoding="utf-8") as f:
        f.write("window.OPEX_DATA = " + payload + ";\n")
    log(f"WROTE {out}  ({os.path.getsize(out)/1024:.0f} KB)")
    return out


def main():
    write_data_js(build_data())
    log("Done. Open dashboard.html.")


if __name__ == "__main__":
    main()
